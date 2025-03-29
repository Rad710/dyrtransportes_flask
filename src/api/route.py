import io

from typing import Optional
from typing import Sequence
from typing import Tuple
from typing import List

from dataclasses import asdict

from flask import request
from flask import jsonify
from flask import Response
from flask import make_response

from sqlalchemy import select
from sqlalchemy import asc
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.exc import OperationalError

from openpyxl import Workbook
from openpyxl.styles import numbers, Border, Side
from openpyxl.utils import get_column_letter

from app_config import logger
from app_config import app
from app_config import db_session
from app_config import RequestWithUser

from decorators.token_required import token_required

from models.route import Route
from utils.locale import get_message

request: RequestWithUser

# Translation dictionaries
MESSAGES = {
    "en": {
        # Error messages
        "route_not_found": "Route not found",
        "transaction_error": "Transaction error",
        "get_routes_error": "Error getting routes",
        "invalid_route_data": "Invalid route data",
        "connection_error": "Connection error",
        "add_route_error": "Error adding route",
        "update_route_error": "Error updating route",
        "delete_route_error": "Error deleting route",
        "create_excel_error": "Error creating Excel file",
        # Success messages
        "route_added": "Route added successfully",
        "route_updated": "Route updated successfully",
        "route_deleted": "Route deleted successfully",
        # Excel headers and labels
        "origin": "Origin",
        "destination": "Destination",
        "price": "Price",
        "payroll_price": "Payroll Price",
        "price_list": "price_list",
    },
    "es": {
        # Error messages
        "route_not_found": "No se encontró la ruta",
        "transaction_error": "Error de transacción",
        "get_routes_error": "Error al obtener rutas",
        "invalid_route_data": "Datos de la Ruta inválidos",
        "connection_error": "problema de conexión",
        "add_route_error": "Error al agregar ruta",
        "update_route_error": "Error al actualizar ruta",
        "delete_route_error": "Error al eliminar ruta",
        "create_excel_error": "Error al crear archivo Excel",
        # Success messages
        "route_added": "Ruta agregada exitosamente",
        "route_updated": "Ruta actualizada exitosamente",
        "route_deleted": "Ruta eliminada exitosamente",
        # Excel headers and labels
        "origin": "Origen",
        "destination": "Destino",
        "price": "Precio",
        "payroll_price": "Precio de Liquidación",
        "price_list": "lista_de_precios",
    },
}


@app.route("/api/route/<int:route_code>", methods=["GET"])
@token_required
def get_route(route_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(Route).where(
            Route.route_code == route_code,
            Route.deleted == False,
            Route.modification_user == request.current_user.user_id,
        )

        route: Optional[Route] = db_session.scalar(stmt)
        if route is None:
            logger.error("fetch table Route, not found")
            return jsonify({"message": get_message(MESSAGES, "route_not_found")}), 404

        logger.info("fetch table Route, found: %s", route.route_code)
        logger.debug("fetch table Route, found: %s", route)

        return jsonify(route), 200

    except SQLAlchemyError as e:
        logger.error("fetch table Route, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500


@app.route("/api/routes", methods=["GET"])
@token_required
def get_route_list() -> Tuple[Response, int]:
    try:
        stmt = (
            select(Route)
            .where(
                Route.deleted == False,
                Route.modification_user == request.current_user.user_id,
            )
            .order_by(asc(Route.origin), asc(Route.destination))
        )

        routes: Sequence[Route] = db_session.scalars(stmt).all()
        logger.info("fetch routes table Route, len: %s", len(routes))
        logger.debug("fetch routes table Route, routes: %s", routes)
        return jsonify(routes), 200

    except SQLAlchemyError as e:
        logger.error("fetch routes table Route, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "get_routes_error")}), 500


@app.route("/api/route", methods=["POST"])
@token_required
def post_route() -> Tuple[Response, int]:
    try:
        # json to db object
        payload = Route(
            **request.get_json(), modification_user=request.current_user.user_id
        )

        # add to database
        logger.debug("insert table Route, payload: %s", payload)
        db_session.add(payload)

        db_session.commit()
        logger.info("inserted table Route, route: %s", payload.route_code)
        return (
            jsonify(
                {**asdict(payload), "message": get_message(MESSAGES, "route_added")}
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("insert table Route, invalid route error: %s", e)
        error_msg = f"{get_message(MESSAGES, 'invalid_route_data')} ({e})"
        return jsonify({"message": error_msg}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("insert table Route, connection error: %s", e)

        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "add_route_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("insert table Route, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "add_route_error")}), 500


@app.route("/api/route/<int:route_code>", methods=["PUT"])
@token_required
def put_route(route_code: int) -> Tuple[Response, int]:
    try:
        # get entry to update
        stmt = select(Route).where(
            Route.route_code == route_code,
            Route.modification_user == request.current_user.user_id,
        )
        entry_to_update: Optional[Route] = db_session.scalar(stmt)

        if entry_to_update is None:
            logger.error("update table Route, route not found")
            return jsonify({"message": get_message(MESSAGES, "route_not_found")}), 404

        # json to db object
        payload = Route(
            **request.get_json(), modification_user=request.current_user.user_id
        )

        entry_to_update.origin = payload.origin
        entry_to_update.destination = payload.destination
        entry_to_update.price = payload.price
        entry_to_update.payroll_price = payload.payroll_price
        entry_to_update.modification_user = payload.modification_user

        logger.info("update table Route, payload: %s", entry_to_update)

        db_session.commit()
        logger.info("updated table Route, route: %s", route_code)
        return (
            jsonify(
                {
                    **asdict(entry_to_update),
                    "message": get_message(MESSAGES, "route_updated"),
                }
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("invalid route: %s", e)
        error_msg = f"{get_message(MESSAGES, 'invalid_route_data')} ({e})"
        return jsonify({"message": error_msg}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("update table Route: connection error %s", e)

        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "update_route_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("update table Route, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "update_route_error")}), 500


@app.route("/api/route/<int:route_code>", methods=["DELETE"])
@token_required
def delete_route(route_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(Route).where(
            Route.route_code == route_code,
            Route.modification_user == request.current_user.user_id,
        )
        existing_entry: Optional[Route] = db_session.scalar(stmt)

        if existing_entry is None:
            logger.error("delete table Route, route not found")
            return jsonify({"message": get_message(MESSAGES, "route_not_found")}), 404

        existing_entry.deleted = True
        existing_entry.modification_user = request.current_user.user_id
        db_session.commit()
        logger.info("delete table Route: route %s", route_code)
        return jsonify({"message": get_message(MESSAGES, "route_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table Route, connection error: %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_route_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table Route, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_route_error")}), 500


@app.route("/api/routes", methods=["DELETE"])
@token_required
def delete_routes() -> Tuple[Response, int]:
    if (request.data is None) or (not request.is_json):
        logger.error("delete routes table Route, route list is empty")
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_route_error")
                    + ": "
                    + get_message(MESSAGES, "route_not_found")
                }
            ),
            404,
        )

    route_list: List[int] = request.get_json()
    logger.debug("delete routes table Route, payload: %s", route_list)

    try:
        for route_code in route_list:
            stmt = select(Route).where(
                Route.route_code == route_code,
                Route.modification_user == request.current_user.user_id,
            )
            route: Optional[Route] = db_session.scalar(stmt)

            if route is None:
                return (
                    jsonify(
                        {
                            "message": get_message(MESSAGES, "delete_route_error")
                            + ": "
                            + get_message(MESSAGES, "route_not_found")
                        }
                    ),
                    404,
                )

            route.deleted = True
            route.modification_user = request.current_user.user_id
            logger.info("delete table Route, route: %s", route_code)

        db_session.commit()
        return jsonify({"message": get_message(MESSAGES, "route_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table Route: connection error %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_route_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table Route, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_route_error")}), 500


@app.route("/api/routes/export-excel", methods=["GET"])
@token_required
def routes_export_excel() -> Tuple[Response, int]:
    try:
        stmt = (
            select(Route)
            .where(
                Route.deleted == False,
                Route.modification_user == request.current_user.user_id,
            )
            .order_by(asc(Route.origin), asc(Route.destination))
        )

        route_list: Sequence[Route] = db_session.scalars(stmt).all()

    except SQLAlchemyError as e:
        logger.error("fetch routes table Route, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "create_excel_error")}), 500

    # Create file
    output = io.BytesIO()
    workbook = Workbook(write_only=False, iso_dates=False)
    sheet = workbook.active

    # Get translated headers
    headers = [
        get_message(MESSAGES, "origin"),
        get_message(MESSAGES, "destination"),
        get_message(MESSAGES, "price"),
        get_message(MESSAGES, "payroll_price"),
    ]
    sheet.append(headers)

    for col_idx in range(1, 5):
        sheet.column_dimensions[get_column_letter(col_idx)].width = 20

    # Border style
    border_style = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Apply border style to each cell in the row
    for cell in sheet[sheet.max_row]:
        cell.border = border_style

    # Add data rows
    for route in route_list:
        row = [route.origin, route.destination, route.price, route.payroll_price]

        sheet.append(row)

        for i in range(3, 5):
            cell = sheet.cell(row=sheet.max_row, column=i)
            cell.number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1  # type: ignore

        # Apply border style to each cell in the row
        for cell in sheet[sheet.max_row]:
            cell.border = border_style

    # Save Excel file to output stream
    workbook.save(output)
    output.seek(0)

    # Create response with Excel file - use translated filename
    filename = f'{get_message(MESSAGES, "price_list")}.xlsx'

    # Create the client response with the Excel file
    response = make_response(output.getvalue())
    response.headers["Content-Type"] = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"

    logger.info("exported Routes excel file")

    return response, 200
