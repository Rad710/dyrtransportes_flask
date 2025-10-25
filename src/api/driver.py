import io

from typing import Optional
from typing import Sequence
from typing import Tuple
from typing import List

from dataclasses import asdict

from flask import request
from flask import jsonify
from flask import Response
from flask import make_response

from sqlalchemy import select
from sqlalchemy import asc
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.exc import OperationalError

from openpyxl import Workbook
from openpyxl.styles import Border
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter

from app_config import logger
from app_config import app
from app_config import db_session
from app_config import RequestWithUser

from decorators.token_required import token_required

from models.driver import Driver
from utils.locale import get_message

request: RequestWithUser

# Translation dictionaries
MESSAGES = {
    "en": {
        # Error messages
        "driver_not_found": "Driver not found",
        "transaction_error": "Transaction error",
        "get_drivers_error": "Error getting drivers",
        "invalid_driver_data": "Invalid driver data",
        "connection_error": "Connection error",
        "add_driver_error": "Error adding driver",
        "update_driver_error": "Error updating driver",
        "delete_driver_error": "Error deleting driver",
        "restore_driver_error": "Error restoring driver",
        "export_error": "Error sending Excel file",
        # Success messages
        "driver_added": "Driver added successfully",
        "driver_updated": "Driver updated successfully",
        "driver_deleted": "Driver deleted successfully",
        "driver_restored": "Driver restored successfully",
        # Excel headers and labels
        "id_number": "ID Number",
        "first_name": "First Name",
        "last_name": "Last Name",
        "truck_plate": "Truck Plate",
        "trailer_plate": "Trailer Plate",
        "driver_list": "driver_list",
    },
    "es": {
        # Error messages
        "driver_not_found": "No se encontró el chofer",
        "transaction_error": "Error de transacción",
        "get_drivers_error": "Error al obtener choferes",
        "invalid_driver_data": "Datos del Chofer inválidos",
        "connection_error": "problema de conexión",
        "add_driver_error": "Error al agregar chofer",
        "update_driver_error": "Error al actualizar chofer",
        "delete_driver_error": "Error al eliminar chofer",
        "restore_driver_error": "Error al restaurar chofer",
        "export_error": "Error al enviar archivo Excel",
        # Success messages
        "driver_added": "Chofer agregado exitosamente",
        "driver_updated": "Chofer actualizado exitosamente",
        "driver_deleted": "Chofer eliminado exitosamente",
        "driver_restored": "Chofer restaurado exitosamente",
        # Excel headers and labels
        "id_number": "C.I.",
        "first_name": "Nombre",
        "last_name": "Apellido",
        "truck_plate": "Chapa Camión",
        "trailer_plate": "Chapa Carreta",
        "driver_list": "nomina_de_choferes",
    },
}


@app.route("/api/driver/<int:driver_code>", methods=["GET"])
@token_required
def get_driver(driver_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(Driver).where(
            Driver.driver_code == driver_code,
            Driver.modification_user == request.current_user.user_id,
        )

        driver: Optional[Driver] = db_session.scalar(stmt)
        if driver is None:
            logger.error("fetch table Driver, not found")
            return jsonify({"message": get_message(MESSAGES, "driver_not_found")}), 404

        logger.info("fetch table Driver, found: %s", driver.driver_code)
        logger.debug("fetch table Driver, found: %s", driver)

        return jsonify(driver), 200

    except SQLAlchemyError as e:
        logger.error("fetch table Driver, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500


@app.route("/api/drivers", methods=["GET"])
@token_required
def get_driver_list() -> Tuple[Response, int]:
    try:
        stmt = (
            select(Driver)
            .where(
                Driver.modification_user == request.current_user.user_id,
            )
            .order_by(
                asc(Driver.deleted), asc(Driver.driver_name), asc(Driver.driver_surname)
            )
        )

        drivers: Sequence[Driver] = db_session.scalars(stmt).all()
        logger.info("fetch drivers table Driver, len: %s", len(drivers))
        logger.debug("fetch drivers table Driver, drivers: %s", drivers)
        return jsonify(drivers), 200

    except SQLAlchemyError as e:
        logger.error("fetch drivers table Driver, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "get_drivers_error")}), 500


@app.route("/api/driver", methods=["POST"])
@token_required
def post_driver() -> Tuple[Response, int]:
    try:
        # json to db object
        payload = Driver(
            **{**request.get_json(), "modification_user": request.current_user.user_id}
        )

        # add to database
        logger.debug("insert table Driver, payload: %s", payload)
        db_session.add(payload)

        db_session.commit()
        logger.info("inserted table Driver, driver: %s", payload.driver_code)
        return (
            jsonify(
                {**asdict(payload), "message": get_message(MESSAGES, "driver_added")}
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("insert table Driver, invalid driver error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_driver_data")}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("insert table Driver, connection error: %s", e)

        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "add_driver_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("insert table Driver, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "add_driver_error")}), 500


@app.route("/api/driver/<int:driver_code>", methods=["PUT"])
@token_required
def put_driver(driver_code: int) -> Tuple[Response, int]:
    try:
        # get entry to update
        stmt = select(Driver).where(
            Driver.driver_code == driver_code,
            Driver.modification_user == request.current_user.user_id,
        )
        entry_to_update: Optional[Driver] = db_session.scalar(stmt)

        if entry_to_update is None:
            logger.error("update table Driver, driver not found")
            return jsonify({"message": get_message(MESSAGES, "driver_not_found")}), 404

        # json to db object
        payload = Driver(
            **{**request.get_json(), "modification_user": request.current_user.user_id}
        )

        entry_to_update.driver_id = payload.driver_id
        entry_to_update.driver_name = payload.driver_name
        entry_to_update.driver_surname = payload.driver_surname
        entry_to_update.truck_plate = payload.truck_plate
        entry_to_update.trailer_plate = payload.trailer_plate
        entry_to_update.modification_user = payload.modification_user

        logger.info("update table Driver, payload: %s", entry_to_update)

        db_session.commit()
        logger.info("updated table Driver, driver: %s", driver_code)
        return (
            jsonify(
                {
                    **asdict(entry_to_update),
                    "message": get_message(MESSAGES, "driver_updated"),
                }
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("invalid driver: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_driver_data")}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("update table Driver: connection error %s", e)

        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "update_driver_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("update table Driver, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "update_driver_error")}), 500


@app.route("/api/driver/<int:driver_code>", methods=["DELETE"])
@token_required
def delete_driver(driver_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(Driver).where(
            Driver.driver_code == driver_code,
            Driver.modification_user == request.current_user.user_id,
        )
        existing_entry: Optional[Driver] = db_session.scalar(stmt)

        if existing_entry is None:
            logger.error("delete table Driver, driver not found")
            return jsonify({"message": get_message(MESSAGES, "driver_not_found")}), 404

        existing_entry.deleted = True
        existing_entry.modification_user = request.current_user.user_id
        db_session.commit()
        logger.info("delete table Driver: driver %s", driver_code)
        return jsonify({"message": get_message(MESSAGES, "driver_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table Driver, connection error: %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_driver_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table Driver, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_driver_error")}), 500


@app.route("/api/drivers", methods=["DELETE"])
@token_required
def delete_drivers() -> Tuple[Response, int]:
    if (request.data is None) or (not request.is_json):
        logger.error("delete drivers table Driver, driver list is empty")
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_driver_error")
                    + ": "
                    + get_message(MESSAGES, "driver_not_found")
                }
            ),
            404,
        )

    driver_list: List[int] = request.get_json()
    logger.debug("delete drivers table Driver, payload: %s", driver_list)

    try:
        for driver_code in driver_list:
            stmt = select(Driver).where(
                Driver.driver_code == driver_code,
                Driver.modification_user == request.current_user.user_id,
            )
            driver: Optional[Driver] = db_session.scalar(stmt)

            if driver is None:
                return (
                    jsonify(
                        {
                            "message": get_message(MESSAGES, "delete_driver_error")
                            + ": "
                            + get_message(MESSAGES, "driver_not_found")
                        }
                    ),
                    404,
                )

            driver.deleted = True
            driver.modification_user = request.current_user.user_id
            logger.info("delete table Driver, driver: %s", driver_code)

        db_session.commit()
        return jsonify({"message": get_message(MESSAGES, "driver_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table Driver: connection error %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_driver_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table Driver, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_driver_error")}), 500


@app.route("/api/driver/<int:driver_code>/restore", methods=["PATCH"])
@token_required
def restore_driver(driver_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(Driver).where(
            Driver.driver_code == driver_code,
            Driver.modification_user == request.current_user.user_id,
        )
        existing_entry: Optional[Driver] = db_session.scalar(stmt)

        if existing_entry is None:
            logger.error("restore table Driver, driver not found")
            return jsonify({"message": get_message(MESSAGES, "driver_not_found")}), 404

        existing_entry.deleted = False
        existing_entry.modification_user = request.current_user.user_id
        db_session.commit()
        logger.info("restore table Driver: driver %s", driver_code)
        return jsonify({"message": get_message(MESSAGES, "driver_restored")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("restore table Driver, connection error: %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "restore_driver_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("restore table Driver, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "restore_driver_error")}), 500


@app.route("/api/drivers/export-excel", methods=["GET"])
@token_required
def drivers_export_excel() -> Tuple[Response, int]:
    try:
        stmt = (
            select(Driver)
            .where(
                Driver.deleted == False,
                Driver.modification_user == request.current_user.user_id,
            )
            .order_by(asc(Driver.driver_name), asc(Driver.driver_surname))
        )

        driver_list: Sequence[Driver] = db_session.scalars(stmt).all()

    except SQLAlchemyError as e:
        logger.error("export Drivers, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "export_error")}), 500

    # Create file
    output = io.BytesIO()
    workbook = Workbook(write_only=False, iso_dates=False)
    sheet = workbook.active

    # Get translated headers
    headers = [
        get_message(MESSAGES, "id_number"),
        get_message(MESSAGES, "first_name"),
        get_message(MESSAGES, "last_name"),
        get_message(MESSAGES, "truck_plate"),
        get_message(MESSAGES, "trailer_plate"),
    ]
    sheet.append(headers)

    for col_idx in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = 20

    # Border style
    border_style = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Apply border style to each cell in the row
    for cell in sheet[sheet.max_row]:
        cell.border = border_style

    # Add data rows
    for driver in driver_list:
        row = [
            driver.driver_id,
            driver.driver_name,
            driver.driver_surname,
            driver.truck_plate,
            driver.trailer_plate,
        ]

        sheet.append(row)

        # Apply border style to each cell in the row
        for cell in sheet[sheet.max_row]:
            cell.border = border_style

    # Save Excel file to output stream
    workbook.save(output)
    output.seek(0)

    # Create response with Excel file - use translated filename
    filename = f'{get_message(MESSAGES, "driver_list")}.xlsx'

    # Create the client response with the Excel file
    response = make_response(output.getvalue())
    response.headers["Content-Type"] = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"

    logger.info("exported Drivers excel file")

    return response, 200
