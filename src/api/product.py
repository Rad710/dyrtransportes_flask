import io

from typing import Optional
from typing import Sequence
from typing import Tuple
from typing import List

from dataclasses import asdict

from flask import Blueprint
from flask import request
from flask import jsonify
from flask import Response
from flask import make_response

from sqlalchemy import select
from sqlalchemy import asc
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.exc import OperationalError

from openpyxl import Workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from app_config import logger
from app_config import db_session
from app_config import RequestWithUser

from decorators.token_required import token_required

from models.product import Product
from utils.locale import get_message

product_bp = Blueprint("product", __name__)

request: RequestWithUser

# Translation dictionaries
MESSAGES = {
    "en": {
        # Error messages
        "product_not_found": "Product not found",
        "transaction_error": "Transaction error",
        "get_products_error": "Error getting products",
        "invalid_product_data": "Invalid product data",
        "connection_error": "Connection error",
        "add_product_error": "Error adding product",
        "update_product_error": "Error updating product",
        "delete_product_error": "Error deleting product",
        # Success messages
        "product_added": "Product added successfully",
        "product_updated": "Product updated successfully",
        "product_deleted": "Product deleted successfully",
        # Excel headers and labels
        "code": "Code",
        "product_name": "Product Name",
        "product_list": "product_list",
    },
    "es": {
        # Error messages
        "product_not_found": "No se encontró el producto",
        "transaction_error": "Error de transacción",
        "get_products_error": "Error al obtener productos",
        "invalid_product_data": "Datos del producto inválidos",
        "connection_error": "problema de conexión",
        "add_product_error": "Error al agregar producto",
        "update_product_error": "Error al actualizar producto",
        "delete_product_error": "Error al eliminar producto",
        # Success messages
        "product_added": "Producto agregado exitosamente",
        "product_updated": "Producto actualizado exitosamente",
        "product_deleted": "Producto eliminado exitosamente",
        # Excel headers and labels
        "code": "Código",
        "product_name": "Nombre de Producto",
        "product_list": "lista_de_productos",
    },
}


@product_bp.route("/api/product/<int:product_code>", methods=["GET"])
@token_required
def get_product(product_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(Product).where(
            Product.product_code == product_code,
            Product.deleted == False,
            Product.modification_user == request.current_user.user_id,
        )

        product: Optional[Product] = db_session.scalar(stmt)
        if product is None:
            logger.error("fetch table Product, not found")
            return jsonify({"message": get_message(MESSAGES, "product_not_found")}), 404

        logger.info("fetch table Product, found: %s", product.product_code)
        logger.debug("fetch table Product, found: %s", product)

        return jsonify(product), 200

    except SQLAlchemyError as e:
        logger.error("fetch table Product, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500


@product_bp.route("/api/products", methods=["GET"])
@token_required
def get_product_list() -> Tuple[Response, int]:
    try:
        stmt = (
            select(Product)
            .where(
                Product.deleted == False,
                Product.modification_user == request.current_user.user_id,
            )
            .order_by(asc(Product.product_name))
        )

        products: Sequence[Product] = db_session.scalars(stmt).all()
        logger.info("fetch products table Product, len: %s", len(products))
        logger.debug("fetch products table Product, products: %s", products)
        return jsonify(products), 200

    except SQLAlchemyError as e:
        logger.error("fetch products table Product, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "get_products_error")}), 500


@product_bp.route("/api/product", methods=["POST"])
@token_required
def post_product() -> Tuple[Response, int]:
    try:
        # json to db object
        payload = Product(
            **{**request.get_json(), "modification_user": request.current_user.user_id}
        )

        # add to database
        logger.debug("insert table Product, payload: %s", payload)
        db_session.add(payload)

        db_session.commit()
        logger.info("inserted table Product, product: %s", payload.product_code)
        return (
            jsonify(
                {**asdict(payload), "message": get_message(MESSAGES, "product_added")}
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("insert table Product, invalid product error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_product_data")}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("insert table Product, connection error: %s", e)

        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "add_product_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("insert table Product, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "add_product_error")}), 500


@product_bp.route("/api/product/<int:product_code>", methods=["PUT"])
@token_required
def put_product(product_code: int) -> Tuple[Response, int]:
    try:
        # get entry to update
        stmt = select(Product).where(
            Product.product_code == product_code,
            Product.modification_user == request.current_user.user_id,
        )
        entry_to_update: Optional[Product] = db_session.scalar(stmt)

        if entry_to_update is None:
            logger.error("update table Product, product not found")
            return jsonify({"message": get_message(MESSAGES, "product_not_found")}), 404

        # json to db object
        payload = Product(
            **{**request.get_json(), "modification_user": request.current_user.user_id}
        )

        entry_to_update.product_name = payload.product_name
        entry_to_update.modification_user = payload.modification_user

        logger.info("update table Product, payload: %s", entry_to_update)

        db_session.commit()
        logger.info("updated table Product, product: %s", product_code)
        return (
            jsonify(
                {
                    **asdict(entry_to_update),
                    "message": get_message(MESSAGES, "product_updated"),
                }
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("invalid product: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_product_data")}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("update table Product: connection error %s", e)

        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "update_product_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("update table Product, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "update_product_error")}), 500


@product_bp.route("/api/product/<int:product_code>", methods=["DELETE"])
@token_required
def delete_product(product_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(Product).where(
            Product.product_code == product_code,
            Product.modification_user == request.current_user.user_id,
        )
        existing_entry: Optional[Product] = db_session.scalar(stmt)

        if existing_entry is None:
            logger.error("delete table Product, product not found")
            return jsonify({"message": get_message(MESSAGES, "product_not_found")}), 404

        existing_entry.deleted = True
        existing_entry.modification_user = request.current_user.user_id
        db_session.commit()
        logger.info("delete table Product: product %s", product_code)
        return jsonify({"message": get_message(MESSAGES, "product_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table Product, connection error: %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_product_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table Product, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_product_error")}), 500


@product_bp.route("/api/products", methods=["DELETE"])
@token_required
def delete_products() -> Tuple[Response, int]:
    if (request.data is None) or (not request.is_json):
        logger.error("delete products table Product, product list is empty")
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_product_error")
                    + ": "
                    + get_message(MESSAGES, "product_not_found")
                }
            ),
            404,
        )

    product_list: List[int] = request.get_json()
    logger.debug("delete products table Product, payload: %s", product_list)

    try:
        for product_code in product_list:
            stmt = select(Product).where(
                Product.product_code == product_code,
                Product.modification_user == request.current_user.user_id,
            )
            product: Optional[Product] = db_session.scalar(stmt)

            if product is None:
                return (
                    jsonify(
                        {
                            "message": get_message(MESSAGES, "delete_product_error")
                            + ": "
                            + get_message(MESSAGES, "product_not_found")
                        }
                    ),
                    404,
                )

            product.deleted = True
            product.modification_user = request.current_user.user_id
            logger.info("delete table Product, product: %s", product_code)

        db_session.commit()
        return jsonify({"message": get_message(MESSAGES, "product_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table Product: connection error %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_product_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table Product, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_product_error")}), 500


@product_bp.route("/api/products/export-excel", methods=["GET"])
@token_required
def products_export_excel() -> Tuple[Response, int]:
    try:
        stmt = (
            select(Product)
            .where(
                Product.deleted == False,
                Product.modification_user == request.current_user.user_id,
            )
            .order_by(asc(Product.product_name))
        )

        product_list: Sequence[Product] = db_session.scalars(stmt).all()

    except SQLAlchemyError as e:
        logger.error("fetch products table Product, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "get_products_error")}), 500

    # Create file
    output = io.BytesIO()
    workbook = Workbook(write_only=False, iso_dates=False)
    sheet = workbook.active

    # Get translated headers
    headers = [get_message(MESSAGES, "code"), get_message(MESSAGES, "product_name")]
    sheet.append(headers)

    for col_idx in range(1, 3):
        sheet.column_dimensions[get_column_letter(col_idx)].width = 20

    # Border style
    border_style = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Apply border style to each cell in the row
    for cell in sheet[sheet.max_row]:
        cell.border = border_style

    # Add data rows
    for product in product_list:
        row = [product.product_code, product.product_name]

        sheet.append(row)

        # Apply border style to each cell in the row
        for cell in sheet[sheet.max_row]:
            cell.border = border_style

    # Save Excel file to output stream
    workbook.save(output)
    output.seek(0)

    # Create response with Excel file - use translated filename
    filename = f'{get_message(MESSAGES, "product_list")}.xlsx'

    # Create the client response with the Excel file
    response = make_response(output.getvalue())
    response.headers["Content-Type"] = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"

    logger.info("exported Products excel file")

    return response, 200
