import io

from typing import Tuple

from datetime import datetime

from flask import request
from flask import jsonify
from flask import Response
from flask import make_response

from sqlalchemy import text
from sqlalchemy.exc import SQLAlchemyError


from openpyxl import Workbook
from openpyxl.styles import Border
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter

from app_config import logger
from app_config import app
from app_config import db_session
from app_config import RequestWithUser

from decorators.token_required import token_required
from utils.locale import get_message

request: RequestWithUser

# Translation dictionaries
MESSAGES = {
    "en": {
        # Error messages
        "missing_date_params": "The start_date and end_date parameters are required",
        "transaction_error": "Transaction error",
        "excel_generation_error": "Error generating Excel file",
        # Excel headers
        "driver_code": "Driver Code",
        "driver_name": "Driver Name",
        "shipment_count": "Shipment Count",
        "total_origin_kg": "Total Origin Kg.",
        "total_destination_kg": "Total Destination Kg.",
        "difference": "Difference",
        "shipment_fees": "Shipment Fees (Gs.)",
        "settlements": "Settlements (Gs.)",
        "expenses_with_receipt": "Expenses with Receipt (Gs.)",
        "expenses_without_receipt": "Expenses without Receipt (Gs.)",
        "total_expenses": "Total Expenses (Gs.)",
        # Excel sheet name
        "statistics": "Statistics",
        # Excel filename
        "statistics_data": "statistics",
    },
    "es": {
        # Error messages
        "missing_date_params": "Se requieren los parámetros start_date y end_date",
        "transaction_error": "Error de transacción",
        "excel_generation_error": "Error al generar archivo Excel",
        # Excel headers
        "driver_code": "Código de Conductor",
        "driver_name": "Nombre de Conductor",
        "shipment_count": "Cantidad de Cargas",
        "total_origin_kg": "Total Kg. Origen",
        "total_destination_kg": "Total Kg. Destino",
        "difference": "Diferencia",
        "shipment_fees": "Fletes (Gs.)",
        "settlements": "Liquidaciones (Gs.)",
        "expenses_with_receipt": "Gastos con Recibo (Gs.)",
        "expenses_without_receipt": "Gastos sin Recibo (Gs.)",
        "total_expenses": "Total Gastos (Gs.)",
        # Excel sheet name
        "statistics": "Estadísticas",
        # Excel filename
        "statistics_data": "estadisticas",
    },
}


@app.route("/api/statistics", methods=["GET"])
@token_required
def get_statistics_data() -> Tuple[Response, int]:
    try:
        start_date_str = request.args.get("start_date")
        end_date_str = request.args.get("end_date")

        start_date = datetime.fromisoformat(start_date_str) if start_date_str else None
        end_date = datetime.fromisoformat(end_date_str) if end_date_str else None

        if not start_date or not end_date:
            return (
                jsonify({"error": get_message(MESSAGES, "missing_date_params")}),
                400,
            )

        params = {
            "start_date": start_date.date(),
            "end_date": end_date.date(),
            "user_id": request.current_user.user_id,
        }
        t = text(
            """
            SELECT
                sts.driver_code,
                d.driver_name,
                sts.shipments,
                sts.total_origin_weight,
                sts.total_destination_weight,
                sts.total_diff,
                sts.total_shipment_payroll,
                sts.total_driver_payroll,
                sts.total_expenses_amount_receipt,
                sts.total_expenses_amount_no_receipt,
                (sts.total_expenses_amount_receipt + sts.total_expenses_amount_no_receipt) AS total_expenses_amount
            FROM
                (
                SELECT
                    s.driver_code,
                    shipments,
                    total_origin_weight,
                    total_destination_weight,
                    total_diff,
                    total_shipment_payroll,
                    total_driver_payroll,
                    COALESCE(ser.total_expenses_amount_receipt, 0) AS total_expenses_amount_receipt,
                    COALESCE(senr.total_expenses_amount_no_receipt, 0) AS total_expenses_amount_no_receipt
                FROM
                    (
                    SELECT
                        s.driver_code,
                        COUNT(s.shipment_code) AS shipments,
                        SUM(s.origin_weight) AS total_origin_weight,
                        SUM(s.destination_weight) AS total_destination_weight,
                        SUM(s.destination_weight) - SUM(s.origin_weight) AS total_diff,
                        SUM(s.price * s.destination_weight) AS total_shipment_payroll,
                        SUM(s.payroll_price * s.destination_weight) AS total_driver_payroll
                    FROM
                        shipment s
                    WHERE
                        CAST(s.shipment_date AS DATE) >= :start_date and CAST(s.shipment_date AS DATE) <= :end_date
                        AND s.modification_user = :user_id
                    GROUP BY
                        s.driver_code
                ) s
                LEFT JOIN (
                    SELECT
                        dp.driver_code,
                        SUM(se.amount) AS total_expenses_amount_receipt
                    FROM
                        shipment_expense se
                    INNER JOIN driver_payroll dp ON
                        dp.payroll_code = se.driver_payroll_code
                    WHERE
                        se.receipt IS NOT NULL
                        AND CAST(se.expense_date AS DATE) >= :start_date and CAST(se.expense_date AS DATE) <= :end_date
                        AND se.modification_user = :user_id
                    GROUP BY
                        dp.driver_code
                ) ser ON
                    ser.driver_code = s.driver_code
                LEFT JOIN (
                    SELECT
                        dp.driver_code,
                        SUM(se.amount) AS total_expenses_amount_no_receipt
                    FROM
                        shipment_expense se
                    INNER JOIN driver_payroll dp ON
                        dp.payroll_code = se.driver_payroll_code
                    WHERE
                        se.receipt IS NULL
                        AND CAST(se.expense_date AS DATE) >= :start_date and CAST(se.expense_date AS DATE) <= :end_date
                        AND se.modification_user = :user_id
                    GROUP BY
                        dp.driver_code
                ) senr ON
                    senr.driver_code = s.driver_code
            ) sts
            INNER JOIN driver d ON
                d.driver_code = sts.driver_code
            """
        )

        # Execute the query
        statistics = db_session.execute(t, params).all()

        logger.debug("fetch statistics, found: %s", statistics)

        return (
            jsonify(
                [
                    {
                        "driver_code": row.driver_code,
                        "driver_name": row.driver_name,
                        "shipments": row.shipments,
                        "total_origin_weight": row.total_origin_weight,
                        "total_destination_weight": row.total_destination_weight,
                        "total_diff": row.total_diff,
                        "total_shipment_payroll": row.total_shipment_payroll,
                        "total_driver_payroll": row.total_driver_payroll,
                        "total_expenses_amount_receipt": row.total_expenses_amount_receipt,
                        "total_expenses_amount_no_receipt": row.total_expenses_amount_no_receipt,
                        "total_expenses_amount": row.total_expenses_amount,
                    }
                    for row in statistics
                ]
            ),
            200,
        )

    except SQLAlchemyError as e:
        logger.error("fetch statistics, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500

    except Exception as e:
        logger.error("fetch statistics, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500


@app.route("/api/statistics/export-excel", methods=["GET"])
@token_required
def export_statistics_excel() -> Tuple[Response, int]:
    try:
        start_date_str = request.args.get("start_date")
        end_date_str = request.args.get("end_date")

        start_date = datetime.fromisoformat(start_date_str) if start_date_str else None
        end_date = datetime.fromisoformat(end_date_str) if end_date_str else None

        if not start_date or not end_date:
            return (
                jsonify({"error": get_message(MESSAGES, "missing_date_params")}),
                400,
            )

        params = {
            "start_date": start_date.date(),
            "end_date": end_date.date(),
            "user_id": request.current_user.user_id,
        }
        t = text(
            """
            SELECT
                sts.driver_code,
                d.driver_name,
                sts.shipments,
                sts.total_origin_weight,
                sts.total_destination_weight,
                sts.total_diff,
                sts.total_shipment_payroll,
                sts.total_driver_payroll,
                sts.total_expenses_amount_receipt,
                sts.total_expenses_amount_no_receipt,
                (sts.total_expenses_amount_receipt + sts.total_expenses_amount_no_receipt) AS total_expenses_amount
            FROM
                (
                SELECT
                    s.driver_code,
                    shipments,
                    total_origin_weight,
                    total_destination_weight,
                    total_diff,
                    total_shipment_payroll,
                    total_driver_payroll,
                    COALESCE(ser.total_expenses_amount_receipt, 0) AS total_expenses_amount_receipt,
                    COALESCE(senr.total_expenses_amount_no_receipt, 0) AS total_expenses_amount_no_receipt
                FROM
                    (
                    SELECT
                        s.driver_code,
                        COUNT(s.shipment_code) AS shipments,
                        SUM(s.origin_weight) AS total_origin_weight,
                        SUM(s.destination_weight) AS total_destination_weight,
                        SUM(s.destination_weight) - SUM(s.origin_weight) AS total_diff,
                        SUM(s.price * s.destination_weight) AS total_shipment_payroll,
                        SUM(s.payroll_price * s.destination_weight) AS total_driver_payroll
                    FROM
                        shipment s
                    WHERE
                        CAST(s.shipment_date AS DATE) >= :start_date and CAST(s.shipment_date AS DATE) <= :end_date
                        AND s.modification_user = :user_id
                    GROUP BY
                        s.driver_code
                ) s
                LEFT JOIN (
                    SELECT
                        dp.driver_code,
                        SUM(se.amount) AS total_expenses_amount_receipt
                    FROM
                        shipment_expense se
                    INNER JOIN driver_payroll dp ON
                        dp.payroll_code = se.driver_payroll_code
                    WHERE
                        se.receipt IS NOT NULL
                        AND CAST(se.expense_date AS DATE) >= :start_date and CAST(se.expense_date AS DATE) <= :end_date
                        AND se.modification_user = :user_id
                    GROUP BY
                        dp.driver_code
                ) ser ON
                    ser.driver_code = s.driver_code
                LEFT JOIN (
                    SELECT
                        dp.driver_code,
                        SUM(se.amount) AS total_expenses_amount_no_receipt
                    FROM
                        shipment_expense se
                    INNER JOIN driver_payroll dp ON
                        dp.payroll_code = se.driver_payroll_code
                    WHERE
                        se.receipt IS NULL
                        AND CAST(se.expense_date AS DATE) >= :start_date and CAST(se.expense_date AS DATE) <= :end_date
                        AND se.modification_user = :user_id
                    GROUP BY
                        dp.driver_code
                ) senr ON
                    senr.driver_code = s.driver_code
            ) sts
            INNER JOIN driver d ON
                d.driver_code = sts.driver_code
            """
        )

        # Execute the query
        statistics = db_session.execute(t, params).all()

        logger.debug("fetch statistics for Excel export, found: %s", statistics)

        # Create Excel file
        output = io.BytesIO()
        workbook = Workbook(write_only=False, iso_dates=False)
        sheet = workbook.active

        # Get translated sheet name
        sheet.title = get_message(MESSAGES, "statistics")

        # Define headers with translations
        headers = [
            get_message(MESSAGES, "driver_code"),
            get_message(MESSAGES, "driver_name"),
            get_message(MESSAGES, "shipment_count"),
            get_message(MESSAGES, "total_origin_kg"),
            get_message(MESSAGES, "total_destination_kg"),
            get_message(MESSAGES, "difference"),
            get_message(MESSAGES, "shipment_fees"),
            get_message(MESSAGES, "settlements"),
            get_message(MESSAGES, "expenses_with_receipt"),
            get_message(MESSAGES, "expenses_without_receipt"),
            get_message(MESSAGES, "total_expenses"),
        ]
        sheet.append(headers)

        # Set column width
        for col_idx in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 25

        # Border style
        border_style = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Apply border to header cells
        for cell in sheet[sheet.max_row]:
            cell.border = border_style

        # Add data rows
        for row in statistics:
            data_row = [
                row.driver_code,
                row.driver_name,
                row.shipments,
                (
                    float(row.total_origin_weight)
                    if row.total_origin_weight is not None
                    else 0
                ),
                (
                    float(row.total_destination_weight)
                    if row.total_destination_weight is not None
                    else 0
                ),
                (float(row.total_diff) if row.total_diff is not None else 0),
                (
                    float(row.total_shipment_payroll)
                    if row.total_shipment_payroll is not None
                    else 0
                ),
                (
                    float(row.total_driver_payroll)
                    if row.total_driver_payroll is not None
                    else 0
                ),
                (
                    float(row.total_expenses_amount_receipt)
                    if row.total_expenses_amount_receipt is not None
                    else 0
                ),
                (
                    float(row.total_expenses_amount_no_receipt)
                    if row.total_expenses_amount_no_receipt is not None
                    else 0
                ),
                (
                    float(row.total_expenses_amount)
                    if row.total_expenses_amount is not None
                    else 0
                ),
            ]

            sheet.append(data_row)

            # Apply border to each cell in the row
            for cell in sheet[sheet.max_row]:
                cell.border = border_style

        # Save Excel file to output stream
        workbook.save(output)
        output.seek(0)

        # Get translated filename component
        filename_base = get_message(MESSAGES, "statistics_data")

        # Create response with Excel file
        response = make_response(output.getvalue())
        response.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response.headers["Content-Disposition"] = (
            f"attachment; filename={filename_base}_{start_date.strftime('%Y%m%d')}_a_{end_date.strftime('%Y%m%d')}.xlsx"
        )

        logger.info("exported statistics data to Excel file")

        return response, 200

    except SQLAlchemyError as e:
        logger.error("export statistics Excel, error: %s", e)
        return (
            jsonify({"message": get_message(MESSAGES, "excel_generation_error")}),
            500,
        )

    except Exception as e:
        logger.error("export statistics Excel, error: %s", e)
        return (
            jsonify({"message": get_message(MESSAGES, "excel_generation_error")}),
            500,
        )
