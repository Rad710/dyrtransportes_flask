import io
from datetime import datetime
from typing import Optional
from typing import Sequence
from typing import Tuple
from typing import List

from dataclasses import asdict

from flask import Blueprint
from flask import request, jsonify, Response, make_response
from sqlalchemy import select
from sqlalchemy import desc
from sqlalchemy import cast
from sqlalchemy import Date
from sqlalchemy.sql import extract
from sqlalchemy.exc import SQLAlchemyError, OperationalError

from openpyxl import Workbook
from openpyxl.styles import Border
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter

from app_config import logger, db_session, RequestWithUser
from decorators.token_required import token_required
from models.shipment_payroll import ShipmentPayroll
from utils.locale import get_message

shipment_payroll_bp = Blueprint("shipment_payroll", __name__)

request: RequestWithUser

# Translation dictionaries
MESSAGES = {
    "en": {
        # Error messages
        "payroll_not_found": "Payroll not found",
        "transaction_error": "Transaction error",
        "get_payrolls_error": "Error getting payrolls",
        "invalid_parameters": "Invalid parameters",
        "invalid_payroll_data": "Invalid payroll data",
        "connection_error": "Connection error",
        "add_payroll_error": "Error adding payroll",
        "update_payroll_error": "Error updating payroll",
        "collection_field_required": "Field 'collected' is required",
        "invalid_data": "Invalid data",
        "update_collection_status_error": "Error updating collection status",
        "delete_payroll_error": "Error deleting payroll",
        "empty_payroll_list": "Empty payroll list",
        "export_error": "Error generating the export file",
        "invalid_date_format": "Invalid date format",
        "no_export_data": "No data to export in the selected date range",
        # Success messages
        "payroll_added": "Payroll added successfully",
        "payroll_updated": "Payroll updated successfully",
        "payroll_deleted": "Payroll deleted successfully",
        "collection_status_updated_to_collected": "Collection status updated to collected",
        "collection_status_updated_to_uncollected": "Collection status updated to uncollected",
        # Excel headers and labels
        "code": "Code",
        "date": "Date",
        "collected": "Collected",
        "collection_date": "Collection Date",
        "payroll_list": "payroll_list",
        "yes": "Yes",
        "no": "No",
    },
    "es": {
        # Error messages
        "payroll_not_found": "No se encontró la planilla",
        "transaction_error": "Error de transacción",
        "get_payrolls_error": "Error al obtener planillas",
        "invalid_parameters": "Parámetros inválidos",
        "invalid_payroll_data": "Datos de la Planilla inválidos",
        "connection_error": "problema de conexión",
        "add_payroll_error": "Error al agregar planilla",
        "update_payroll_error": "Error al actualizar planilla",
        "collection_field_required": "Campo 'collected' requerido",
        "invalid_data": "Datos inválidos",
        "update_collection_status_error": "Error al actualizar estado de cobranza",
        "delete_payroll_error": "Error al eliminar planilla",
        "empty_payroll_list": "planilla no encontrada",
        "export_error": "Error al generar el archivo de exportación",
        "invalid_date_format": "Formato de fecha inválido",
        "no_export_data": "No hay datos para exportar en el rango de fechas seleccionado",
        # Success messages
        "payroll_added": "Planilla agregada exitosamente",
        "payroll_updated": "Planilla actualizada exitosamente",
        "payroll_deleted": "Planilla eliminada exitosamente",
        "collection_status_updated_to_collected": "Estado de cobranza actualizado a cobrado",
        "collection_status_updated_to_uncollected": "Estado de cobranza actualizado a no cobrado",
        # Excel headers and labels
        "code": "Código",
        "date": "Fecha",
        "collected": "Cobrado",
        "collection_date": "Fecha de Cobro",
        "payroll_list": "lista_de_planillas",
        "yes": "Sí",
        "no": "No",
    },
}


@shipment_payroll_bp.route("/api/shipment-payroll/<int:payroll_code>", methods=["GET"])
@token_required
def get_shipment_payroll(payroll_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(ShipmentPayroll).where(
            ShipmentPayroll.payroll_code == payroll_code,
            ShipmentPayroll.deleted == False,
            ShipmentPayroll.modification_user == request.current_user.user_id,
        )

        shipment_payroll: Optional[ShipmentPayroll] = db_session.scalar(stmt)
        if shipment_payroll is None:
            logger.error("fetch table ShipmentPayroll, not found")
            return jsonify({"message": get_message(MESSAGES, "payroll_not_found")}), 404

        logger.info(
            "fetch table ShipmentPayroll, found: %s", shipment_payroll.payroll_code
        )
        logger.debug("fetch table ShipmentPayroll, found: %s", shipment_payroll)

        return jsonify(shipment_payroll), 200

    except SQLAlchemyError as e:
        logger.error("fetch table ShipmentPayroll, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500


@shipment_payroll_bp.route("/api/shipment-payrolls", methods=["GET"])
@token_required
def get_shipment_payroll_list() -> Tuple[Response, int]:
    year_param: str | None = request.args.get("year")
    try:
        year: Optional[int] = int(year_param) if year_param else None
    except ValueError as e:
        logger.error("Invalid 'year' parameter %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_parameters")}), 400

    try:
        stmt = select(ShipmentPayroll).where(
            ShipmentPayroll.deleted == False,
            ShipmentPayroll.modification_user == request.current_user.user_id,
        )

        if year:
            stmt = stmt.where(
                extract("year", ShipmentPayroll.payroll_timestamp) == year
            )

        stmt = stmt.order_by(
            desc(ShipmentPayroll.payroll_timestamp), desc(ShipmentPayroll.payroll_code)
        )

        shipment_payrolls: Sequence[ShipmentPayroll] = db_session.scalars(stmt).all()
        logger.info(
            "fetch shipment payrolls table ShipmentPayroll, len: %s",
            len(shipment_payrolls),
        )
        logger.debug(
            "fetch shipment payrolls table ShipmentPayroll, payrolls: %s",
            shipment_payrolls,
        )
        return jsonify(shipment_payrolls), 200

    except SQLAlchemyError as e:
        logger.error("fetch shipment payrolls table ShipmentPayroll, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "get_payrolls_error")}), 500


@shipment_payroll_bp.route("/api/shipment-payroll", methods=["POST"])
@token_required
def post_shipment_payroll() -> Tuple[Response, int]:
    try:
        # json to db object
        payload = ShipmentPayroll(
            **{**request.get_json(), "modification_user": request.current_user.user_id}
        )

        # add to database
        logger.debug("insert table ShipmentPayroll, payload: %s", payload)
        db_session.add(payload)

        db_session.commit()
        logger.info("inserted table ShipmentPayroll, payroll: %s", payload.payroll_code)
        return (
            jsonify(
                {**asdict(payload), "message": get_message(MESSAGES, "payroll_added")}
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("insert table ShipmentPayroll, invalid payroll error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_payroll_data")}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("insert table ShipmentPayroll, connection error: %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "add_payroll_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("insert table ShipmentPayroll, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "add_payroll_error")}), 500


@shipment_payroll_bp.route("/api/shipment-payroll/<int:payroll_code>", methods=["PUT"])
@token_required
def put_shipment_payroll(payroll_code: int) -> Tuple[Response, int]:
    try:
        # get entry to update
        stmt = select(ShipmentPayroll).where(
            ShipmentPayroll.payroll_code == payroll_code,
            ShipmentPayroll.modification_user == request.current_user.user_id,
        )
        entry_to_update: Optional[ShipmentPayroll] = db_session.scalar(stmt)

        if entry_to_update is None:
            logger.error("update table ShipmentPayroll, payroll not found")
            return jsonify({"message": get_message(MESSAGES, "payroll_not_found")}), 404

        # json to db object
        payload = ShipmentPayroll(
            **{**request.get_json(), "modification_user": request.current_user.user_id}
        )

        entry_to_update.payroll_timestamp = payload.payroll_timestamp
        entry_to_update.collected = payload.collected
        entry_to_update.collection_timestamp = payload.collection_timestamp
        entry_to_update.deleted = payload.deleted
        entry_to_update.modification_user = payload.modification_user

        logger.info("update table ShipmentPayroll, payload: %s", entry_to_update)

        db_session.commit()
        logger.info("updated table ShipmentPayroll, payroll: %s", payroll_code)
        return (
            jsonify(
                {
                    **asdict(entry_to_update),
                    "message": get_message(MESSAGES, "payroll_updated"),
                }
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("invalid payroll: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_payroll_data")}), 500

    except OperationalError as e:
        db_session.rollback()
        logger.error("update table ShipmentPayroll: connection error %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "update_payroll_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("update table ShipmentPayroll, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "update_payroll_error")}), 500


@shipment_payroll_bp.route(
    "/api/shipment-payroll/<int:payroll_code>/collection-status", methods=["PATCH"]
)
@token_required
def update_shipment_payroll_collection_status(
    payroll_code: int,
) -> Tuple[Response, int]:
    try:
        # Get the entry to update
        stmt = select(ShipmentPayroll).where(
            ShipmentPayroll.payroll_code == payroll_code,
            ShipmentPayroll.modification_user == request.current_user.user_id,
        )
        entry_to_update: Optional[ShipmentPayroll] = db_session.scalar(stmt)

        if entry_to_update is None:
            logger.error(
                "update collection status, payroll not found: %s", payroll_code
            )
            return jsonify({"message": get_message(MESSAGES, "payroll_not_found")}), 404

        # Get the payload
        payload = request.get_json()

        if "collected" not in payload:
            logger.error("update collection status, missing collected field")
            return (
                jsonify(
                    {"message": get_message(MESSAGES, "collection_field_required")}
                ),
                400,
            )

        # Update only the collection-related fields
        entry_to_update.collected = payload["collected"]
        # Set collection_timestamp to current timestamp if collected, otherwise set to None
        if payload["collected"]:
            entry_to_update.collection_timestamp = datetime.now()
        else:
            entry_to_update.collection_timestamp = None

        entry_to_update.modification_user = request.current_user.user_id

        logger.info(
            "update collection status, payroll: %s, status: %s",
            payroll_code,
            "collected" if entry_to_update.collected else "uncollected",
        )

        db_session.commit()

        status_message = get_message(
            MESSAGES,
            (
                "collection_status_updated_to_collected"
                if entry_to_update.collected
                else "collection_status_updated_to_uncollected"
            ),
        )

        return (
            jsonify(
                {
                    **asdict(entry_to_update),
                    "message": status_message,
                }
            ),
            200,
        )

    except (TypeError, ValueError, KeyError) as e:
        logger.error("invalid collection status update: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_data")}), 400

    except OperationalError as e:
        db_session.rollback()
        logger.error("update collection status: connection error %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "update_collection_status_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("update collection status, error: %s", e)
        return (
            jsonify(
                {"message": get_message(MESSAGES, "update_collection_status_error")}
            ),
            500,
        )


@shipment_payroll_bp.route("/api/shipment-payroll/<int:payroll_code>", methods=["DELETE"])
@token_required
def delete_shipment_payroll(payroll_code: int) -> Tuple[Response, int]:
    try:
        stmt = select(ShipmentPayroll).where(
            ShipmentPayroll.payroll_code == payroll_code,
            ShipmentPayroll.modification_user == request.current_user.user_id,
        )
        existing_entry: Optional[ShipmentPayroll] = db_session.scalar(stmt)

        if existing_entry is None:
            logger.error("delete table ShipmentPayroll, payroll not found")
            return jsonify({"message": get_message(MESSAGES, "payroll_not_found")}), 404

        existing_entry.deleted = True
        existing_entry.modification_user = request.current_user.user_id
        db_session.commit()
        logger.info("delete table ShipmentPayroll: payroll %s", payroll_code)
        return jsonify({"message": get_message(MESSAGES, "payroll_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table ShipmentPayroll, connection error: %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_payroll_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table ShipmentPayroll, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_payroll_error")}), 500


@shipment_payroll_bp.route("/api/shipment-payrolls", methods=["DELETE"])
@token_required
def delete_shipment_payrolls() -> Tuple[Response, int]:
    if (request.data is None) or (not request.is_json):
        logger.error("delete payrolls table ShipmentPayroll, payroll list is empty")
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_payroll_error")
                    + ": "
                    + get_message(MESSAGES, "empty_payroll_list")
                }
            ),
            404,
        )

    payroll_list: List[int] = request.get_json()
    logger.debug("delete payrolls table ShipmentPayroll, payload: %s", payroll_list)

    try:
        for payroll_code in payroll_list:
            stmt = select(ShipmentPayroll).where(
                ShipmentPayroll.payroll_code == payroll_code,
                ShipmentPayroll.modification_user == request.current_user.user_id,
            )
            payroll: Optional[ShipmentPayroll] = db_session.scalar(stmt)

            if payroll is None:
                return (
                    jsonify(
                        {
                            "message": get_message(MESSAGES, "delete_payroll_error")
                            + ": "
                            + get_message(MESSAGES, "empty_payroll_list")
                        }
                    ),
                    404,
                )

            payroll.deleted = True
            payroll.modification_user = request.current_user.user_id
            logger.info("delete table ShipmentPayroll, payroll: %s", payroll_code)

        db_session.commit()
        return jsonify({"message": get_message(MESSAGES, "payroll_deleted")}), 200

    except OperationalError as e:
        db_session.rollback()
        logger.error("delete table ShipmentPayroll: connection error %s", e)
        return (
            jsonify(
                {
                    "message": get_message(MESSAGES, "delete_payroll_error")
                    + ": "
                    + get_message(MESSAGES, "connection_error")
                }
            ),
            503,
        )

    except SQLAlchemyError as e:
        db_session.rollback()
        logger.error("delete table ShipmentPayroll, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "delete_payroll_error")}), 500


@shipment_payroll_bp.route("/api/shipment-payrolls/export-excel", methods=["GET"])
@token_required
def export_shipment_payrolls() -> Tuple[Response, int]:
    # Get date range parameters
    start_date_str = request.args.get("start_date", "")
    end_date_str = request.args.get("end_date", "")

    try:
        # Convert date parameters as needed
        start_date = datetime.fromisoformat(start_date_str) if start_date_str else None
        end_date = datetime.fromisoformat(end_date_str) if end_date_str else None

        # Build query for shipment payrolls in date range
        stmt = select(ShipmentPayroll).where(
            ShipmentPayroll.deleted == False,
            ShipmentPayroll.modification_user == request.current_user.user_id,
        )

        if start_date:
            stmt = stmt.where(
                cast(ShipmentPayroll.payroll_timestamp, Date) >= start_date.date()
            )
        if end_date:
            stmt = stmt.where(
                cast(ShipmentPayroll.payroll_timestamp, Date) <= end_date.date()
            )

        stmt = stmt.order_by(
            desc(ShipmentPayroll.payroll_timestamp), desc(ShipmentPayroll.payroll_code)
        )

        payroll_list: Sequence[ShipmentPayroll] = db_session.scalars(stmt).all()

        if not payroll_list:
            logger.info("export ShipmentPayrolls, no data found in date range")
            return (
                jsonify({"message": get_message(MESSAGES, "no_export_data")}),
                404,
            )

        # Create file
        output = io.BytesIO()
        workbook = Workbook(write_only=False, iso_dates=False)
        sheet = workbook.active

        # Get translated headers
        headers = [
            get_message(MESSAGES, "code"),
            get_message(MESSAGES, "date"),
            get_message(MESSAGES, "collected"),
            get_message(MESSAGES, "collection_date"),
        ]
        sheet.append(headers)

        for col_idx in range(1, 5):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 20

        # Border style
        border_style = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Apply border style to each cell in the header row
        for cell in sheet[sheet.max_row]:
            cell.border = border_style

        # Add data rows
        for payroll in payroll_list:
            # Format dates for display
            payroll_date = (
                payroll.payroll_timestamp.strftime("%d/%m/%Y")
                if payroll.payroll_timestamp
                else ""
            )
            collection_date = (
                payroll.collection_timestamp.strftime("%d/%m/%Y")
                if payroll.collection_timestamp
                else ""
            )

            # Get translated yes/no values
            collected_text = (
                get_message(MESSAGES, "yes")
                if payroll.collected
                else get_message(MESSAGES, "no")
            )

            row = [
                payroll.payroll_code,
                payroll_date,
                collected_text,
                collection_date,
            ]

            sheet.append(row)

            # Apply border style to each cell in the data row
            for cell in sheet[sheet.max_row]:
                cell.border = border_style

        # Save Excel file to output stream
        workbook.save(output)
        output.seek(0)

        # Get translated filename
        filename = f'{get_message(MESSAGES, "payroll_list")}.xlsx'

        # Create response with Excel file
        response = make_response(output.getvalue())
        response.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response.headers["Content-Disposition"] = f"attachment; filename={filename}"

        logger.info("exported ShipmentPayrolls excel file")

        return response, 200

    except ValueError as e:
        logger.error("export ShipmentPayrolls, invalid date format: %s", e)
        return jsonify({"message": get_message(MESSAGES, "invalid_date_format")}), 400

    except SQLAlchemyError as e:
        logger.error("export ShipmentPayrolls, database error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "export_error")}), 500

    except Exception as e:
        logger.error("export ShipmentPayrolls, unexpected error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "export_error")}), 500
