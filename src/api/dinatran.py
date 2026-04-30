import io

from typing import Tuple

from datetime import datetime

from flask import Blueprint
from flask import request
from flask import jsonify
from flask import Response
from flask import make_response

from sqlalchemy import func
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy import cast
from sqlalchemy import Date

from openpyxl import Workbook
from openpyxl.styles import Border
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter

from app_config import logger
from app_config import db_session
from app_config import RequestWithUser

from models.shipment import Shipment

from decorators.token_required import token_required
from utils.locale import get_message

dinatran_bp = Blueprint("dinatran", __name__)

request: RequestWithUser

# Translation dictionaries
MESSAGES = {
    "en": {
        "missing_date_params": "The start_date and end_date parameters are required",
        "transaction_error": "Transaction error",
        "excel_generation_error": "Error generating Excel file",
        # Excel headers
        "plate": "Plate",
        "shipment_count": "Shipment Count",
        "total_origin_kg": "Total Origin Kg.",
        "total_destination_kg": "Total Destination Kg.",
        "shipment_fees": "Shipment Fees (Gs.)",
        "settlements": "Settlements (Gs.)",
        # Filename
        "dinatran_data": "dinatran_data",
    },
    "es": {
        "missing_date_params": "Se requieren los parámetros start_date y end_date",
        "transaction_error": "Error de transacción",
        "excel_generation_error": "Error al generar archivo Excel",
        # Excel headers
        "plate": "Chapa",
        "shipment_count": "Cantidad de Cargas",
        "total_origin_kg": "Total Kg. Origen",
        "total_destination_kg": "Total Kg. Destino",
        "shipment_fees": "Fletes (Gs.)",
        "settlements": "Liquidaciones (Gs.)",
        # Filename
        "dinatran_data": "datos_dinatran",
    },
}


@dinatran_bp.route("/api/dinatran", methods=["GET"])
@token_required
def get_dinatran_data() -> Tuple[Response, int]:
    try:
        start_date_str = request.args.get("start_date")
        end_date_str = request.args.get("end_date")

        start_date = datetime.fromisoformat(start_date_str) if start_date_str else None
        end_date = datetime.fromisoformat(end_date_str) if end_date_str else None

        if not start_date or not end_date:
            return (
                jsonify({"error": get_message(MESSAGES, "missing_date_params")}),
                400,
            )

        stmt = (
            select(
                Shipment.truck_plate.label("truck_plate"),
                func.count(Shipment.shipment_code).label("shipments"),
                func.sum(Shipment.origin_weight).label("total_origin_weight"),
                func.sum(Shipment.destination_weight).label("total_destination_weight"),
                (
                    func.sum(Shipment.destination_weight)
                    - func.sum(Shipment.origin_weight)
                ).label("total_diff"),
                func.sum(Shipment.price * Shipment.destination_weight).label(
                    "total_shipment_payroll"
                ),
                func.sum(Shipment.payroll_price * Shipment.destination_weight).label(
                    "total_driver_payroll"
                ),
            )
            .where(
                cast(Shipment.shipment_date, Date) >= start_date.date(),
                cast(Shipment.shipment_date, Date) <= end_date.date(),
                Shipment.modification_user == request.current_user.user_id,
            )
            .group_by(Shipment.truck_plate)
            .order_by(Shipment.truck_plate)
        )

        # Execute the query
        dinatran_shipments_grouped = db_session.execute(stmt).all()

        logger.debug("fetch DINATRAN, found: %s", dinatran_shipments_grouped)

        return (
            jsonify(
                [
                    {
                        "truck_plate": row.truck_plate,
                        "shipments": row.shipments,
                        "total_origin_weight": row.total_origin_weight,
                        "total_destination_weight": row.total_destination_weight,
                        "total_diff": row.total_diff,
                        "total_shipment_payroll": row.total_shipment_payroll,
                        "total_driver_payroll": row.total_driver_payroll,
                    }
                    for row in dinatran_shipments_grouped
                ]
            ),
            200,
        )

    except SQLAlchemyError as e:
        logger.error("fetch DINATRAN, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500

    except Exception as e:
        logger.error("fetch DINATRAN, error: %s", e)
        return jsonify({"message": get_message(MESSAGES, "transaction_error")}), 500


@dinatran_bp.route("/api/dinatran/export-excel", methods=["GET"])
@token_required
def export_dinatran_excel() -> Tuple[Response, int]:
    try:
        start_date_str = request.args.get("start_date")
        end_date_str = request.args.get("end_date")

        start_date = datetime.fromisoformat(start_date_str) if start_date_str else None
        end_date = datetime.fromisoformat(end_date_str) if end_date_str else None

        if not start_date or not end_date:
            return (
                jsonify({"error": get_message(MESSAGES, "missing_date_params")}),
                400,
            )

        stmt = (
            select(
                Shipment.truck_plate.label("truck_plate"),
                func.count(Shipment.shipment_code).label("shipments"),
                func.sum(Shipment.origin_weight).label("total_origin_weight"),
                func.sum(Shipment.destination_weight).label("total_destination_weight"),
                func.sum(Shipment.price * Shipment.destination_weight).label(
                    "total_shipment_payroll"
                ),
                func.sum(Shipment.payroll_price * Shipment.destination_weight).label(
                    "total_driver_payroll"
                ),
            )
            .where(
                cast(Shipment.shipment_date, Date) >= start_date.date(),
                cast(Shipment.shipment_date, Date) <= end_date.date(),
                Shipment.modification_user == request.current_user.user_id,
            )
            .group_by(Shipment.truck_plate)
            .order_by(Shipment.truck_plate)
        )

        # Execute the query
        dinatran_shipments_grouped = db_session.execute(stmt).all()

        logger.debug(
            "fetch DINATRAN for Excel export, found: %s", dinatran_shipments_grouped
        )

        # Create Excel file
        output = io.BytesIO()
        workbook = Workbook(write_only=False, iso_dates=False)
        sheet = workbook.active
        sheet.title = "DINATRAN"

        # Define headers with translations
        headers = [
            get_message(MESSAGES, "plate"),
            get_message(MESSAGES, "shipment_count"),
            get_message(MESSAGES, "total_origin_kg"),
            get_message(MESSAGES, "total_destination_kg"),
            get_message(MESSAGES, "shipment_fees"),
            get_message(MESSAGES, "settlements"),
        ]
        sheet.append(headers)

        # Set column width
        for col_idx in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = 25

        # Border style
        border_style = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Apply border to header cells
        for cell in sheet[sheet.max_row]:
            cell.border = border_style

        # Add data rows
        for row in dinatran_shipments_grouped:
            data_row = [
                row.truck_plate,
                row.shipments,
                (
                    float(row.total_origin_weight)
                    if row.total_origin_weight is not None
                    else 0
                ),
                (
                    float(row.total_destination_weight)
                    if row.total_destination_weight is not None
                    else 0
                ),
                (
                    float(row.total_shipment_payroll)
                    if row.total_shipment_payroll is not None
                    else 0
                ),
                (
                    float(row.total_driver_payroll)
                    if row.total_driver_payroll is not None
                    else 0
                ),
            ]

            sheet.append(data_row)

            # Apply border to each cell in the row
            for cell in sheet[sheet.max_row]:
                cell.border = border_style

        # Save Excel file to output stream
        workbook.save(output)
        output.seek(0)

        # Create response with Excel file using translated filename
        filename = f"{get_message(MESSAGES, 'dinatran_data')}_{start_date.strftime('%Y%m%d')}_a_{end_date.strftime('%Y%m%d')}.xlsx"

        response = make_response(output.getvalue())
        response.headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response.headers["Content-Disposition"] = f"attachment; filename={filename}"

        logger.info("exported DINATRAN data to Excel file")

        return response, 200

    except SQLAlchemyError as e:
        logger.error("export DINATRAN Excel, error: %s", e)
        return (
            jsonify({"message": get_message(MESSAGES, "excel_generation_error")}),
            500,
        )

    except Exception as e:
        logger.error("export DINATRAN Excel, error: %s", e)
        return (
            jsonify({"message": get_message(MESSAGES, "excel_generation_error")}),
            500,
        )
