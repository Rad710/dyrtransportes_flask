"""The PDF exports: the Excel file converted by LibreOffice.

These tests need LibreOffice installed, they are skipped without it.
"""

import pytest

from utils.pdf import is_pdf_conversion_available

pytestmark = [
    pytest.mark.db,
    pytest.mark.pdf,
    pytest.mark.skipif(
        not is_pdf_conversion_available(),
        reason="LibreOffice is not installed, install libreoffice-calc-nogui",
    ),
]


def pdf_text(data: bytes) -> str:
    """Text of the first page, needs pypdfium2 which is a test only dependency."""
    pdfium = pytest.importorskip("pypdfium2")
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".pdf") as pdf_file:
        pdf_file.write(data)
        pdf_file.flush()
        document = pdfium.PdfDocument(pdf_file.name)
        return document[0].get_textpage().get_text_range()


def test_the_collection_pdf_is_a_pdf(client, auth, full_payroll):
    response = client.get(
        f"/api/shipments/export-pdf?shipment_payroll_code={full_payroll['shipment_payroll_code']}",
        headers=auth,
    )

    assert response.status_code == 200
    assert response.headers["Content-Type"] == "application/pdf"
    assert "attachment" in response.headers["Content-Disposition"]
    assert response.data[:5] == b"%PDF-"


def test_the_settlement_pdf_is_a_pdf(client, auth, full_payroll):
    response = client.get(
        f"/api/driver-payroll/export-pdf/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    assert response.status_code == 200
    assert response.headers["Content-Type"] == "application/pdf"
    assert response.data[:5] == b"%PDF-"
    assert response.headers["Content-Disposition"].endswith(".pdf")


def test_the_settlement_pdf_shows_the_dispatch_number(client, auth, full_payroll):
    response = client.get(
        f"/api/driver-payroll/export-pdf/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    text = pdf_text(response.data)

    assert "REM-1000" in text
    assert "LIQUIDACION DE FLETES" in text


def test_libreoffice_calculates_the_formulas_of_the_settlement(
    client, auth, full_payroll
):
    """29.950 kilos at 55,50 per kilo is 1.662.225, the PDF must show the result
    and not the formula."""
    response = client.get(
        f"/api/driver-payroll/export-pdf/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    text = pdf_text(response.data)

    assert "1.662.225" in text
    assert "ROUND(" not in text


def test_the_settlement_pdf_totals_are_calculated(client, auth, full_payroll):
    """One shipment of 1.662.225 and one expense of 100.000 with receipt."""
    response = client.get(
        f"/api/driver-payroll/export-pdf/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    text = pdf_text(response.data).replace("\r\n", " ")

    assert "1.562.225" in text, "total to collect, shipments minus expenses"


def test_the_collection_pdf_calculates_its_total(client, auth, full_payroll):
    """29.950 kilos at 120,75 is 3.616.463, rounded the way Excel rounds."""
    response = client.get(
        f"/api/shipments/export-pdf?shipment_payroll_code={full_payroll['shipment_payroll_code']}",
        headers=auth,
    )

    text = pdf_text(response.data)

    assert "3.616.463" in text


def test_the_numbers_follow_the_language_of_the_request(client, auth, full_payroll):
    spanish = client.get(
        f"/api/shipments/export-pdf?shipment_payroll_code={full_payroll['shipment_payroll_code']}",
        headers={**auth, "Accept-Language": "es-ES,es"},
    )
    english = client.get(
        f"/api/shipments/export-pdf?shipment_payroll_code={full_payroll['shipment_payroll_code']}",
        headers={**auth, "Accept-Language": "en-US,en"},
    )

    assert "3.616.463" in pdf_text(spanish.data)
    assert "3,616,463" in pdf_text(english.data)


def test_a_settlement_pdf_that_does_not_exist_is_not_found(client, auth):
    response = client.get("/api/driver-payroll/export-pdf/999999", headers=auth)

    assert response.status_code == 404


def test_the_pdf_export_needs_a_token(client, full_payroll):
    response = client.get("/api/shipments/export-pdf?shipment_payroll_code=1")

    assert response.status_code == 401


def test_the_pdf_and_the_excel_come_from_the_same_workbook(client, auth, full_payroll):
    """Both endpoints build the same file, so the PDF cannot drift from the Excel."""
    excel = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )
    pdf = client.get(
        f"/api/driver-payroll/export-pdf/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    import io
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(excel.data)).active
    headers = [cell.value for cell in sheet[5] if cell.value]
    text = pdf_text(pdf.data)

    for header in ["Remision N°", "Origen", "Destino"]:
        assert header in headers
        assert header.split()[0] in text
