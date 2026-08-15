"""The Excel and PDF exports, end to end through the API."""

import io
import re
import zipfile

import pytest

from openpyxl import load_workbook

from helpers import http_date

pytestmark = pytest.mark.db

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_of(response):
    return load_workbook(io.BytesIO(response.data))


def sheet_xml_of(response) -> str:
    archive = zipfile.ZipFile(io.BytesIO(response.data))
    return archive.read("xl/worksheets/sheet1.xml").decode()


def cell_values(sheet):
    return [
        [cell.value for cell in row]
        for row in sheet.iter_rows(max_row=sheet.max_row, max_col=sheet.max_column)
    ]


# ---------------------------------------------------------------- collection sheet


def test_the_collection_excel_is_a_spreadsheet(client, auth, full_payroll):
    response = client.get(
        f"/api/shipments/export-excel?shipment_payroll_code={full_payroll['shipment_payroll_code']}",
        headers=auth,
    )

    assert response.status_code == 200
    assert response.headers["Content-Type"] == EXCEL_MIME
    assert "attachment" in response.headers["Content-Disposition"]
    assert response.data[:2] == b"PK"


def test_the_collection_excel_carries_the_shipment(client, auth, full_payroll):
    response = client.get(
        f"/api/shipments/export-excel?shipment_payroll_code={full_payroll['shipment_payroll_code']}",
        headers=auth,
    )

    values = cell_values(workbook_of(response).active)
    flat = [str(value) for row in values for value in row if value is not None]

    assert "REM-1000" in flat, "the dispatch number must be in the sheet"
    assert "REC-2000" in flat
    assert "JUAN PEREZ" in flat


def test_the_collection_excel_is_set_up_for_printing(client, auth, full_payroll):
    response = client.get(
        f"/api/shipments/export-excel?shipment_payroll_code={full_payroll['shipment_payroll_code']}",
        headers=auth,
    )

    sheet = workbook_of(response).active

    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.print_title_rows == "$6:$6"


def test_exporting_a_payroll_with_no_shipments_fails(client, auth, api):
    empty_payroll = api.shipment_payroll(payroll_timestamp=http_date(9))

    response = client.get(
        f"/api/shipments/export-excel?shipment_payroll_code={empty_payroll}",
        headers=auth,
    )

    assert response.status_code >= 400


def test_exporting_a_payroll_that_does_not_exist_is_not_found(client, auth):
    response = client.get(
        "/api/shipments/export-excel?shipment_payroll_code=999999", headers=auth
    )

    assert response.status_code == 404


def test_the_export_rejects_a_payroll_code_that_is_not_a_number(client, auth):
    response = client.get(
        "/api/shipments/export-excel?shipment_payroll_code=abc", headers=auth
    )

    assert response.status_code == 400


def test_the_export_needs_a_token(client, full_payroll):
    response = client.get("/api/shipments/export-excel?shipment_payroll_code=1")

    assert response.status_code == 401


# --------------------------------------------------------------------- settlement


def test_the_settlement_excel_is_a_spreadsheet(client, auth, full_payroll):
    response = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    assert response.status_code == 200
    assert response.headers["Content-Type"] == EXCEL_MIME
    assert response.data[:2] == b"PK"


def test_the_settlement_has_the_dispatch_number_column(client, auth, full_payroll):
    """The column the users asked for: 'que salga el numero de remision'."""
    response = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    sheet = workbook_of(response).active
    headers = [cell.value for cell in sheet[5]]

    assert "Remision N°" in headers
    assert headers.index("Remision N°") < headers.index("Recepcion N°")


def test_the_settlement_lists_the_dispatch_number_of_each_shipment(
    client, auth, full_payroll
):
    response = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    sheet = workbook_of(response).active
    headers = [cell.value for cell in sheet[5]]
    column = headers.index("Remision N°") + 1

    assert sheet.cell(row=6, column=column).value == "REM-1000"


def test_the_settlement_formulas_point_at_the_right_columns(client, auth, full_payroll):
    """The columns moved when the dispatch number was added, the formulas
    have to follow: the amount multiplies arrival kilos by the price."""
    response = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    sheet = workbook_of(response).active
    headers = [cell.value for cell in sheet[5]]
    amount_column = headers.index("Importe Gs.") + 1

    formula = sheet.cell(row=6, column=amount_column).value
    assert formula == "=ROUND($I6*$K6, 0)"
    assert sheet.cell(row=6, column=9).value == 29950  # I, arrival kilos
    assert float(sheet.cell(row=6, column=11).value) == 55.50  # K, price per kilo


def test_the_settlement_is_set_up_for_printing(client, auth, full_payroll):
    response = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    sheet = workbook_of(response).active

    assert sheet.page_setup.orientation == "landscape"
    assert sheet.print_title_rows == "$1:$5"


def test_the_settlement_filename_carries_the_driver_and_the_date(
    client, auth, full_payroll
):
    response = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers=auth,
    )

    disposition = response.headers["Content-Disposition"]

    assert "JUAN PEREZ" in disposition
    assert "Liquidacion" in disposition
    assert disposition.endswith(".xlsx")


def test_a_settlement_that_does_not_exist_is_not_found(client, auth):
    response = client.get("/api/driver-payroll/export-excel/999999", headers=auth)

    assert response.status_code == 404


def test_a_settlement_of_another_user_is_not_exported(
    client, credentials, full_payroll
):
    other = {**credentials, "email": "otro-export@dyrtransportes.com"}
    client.post("/api/auth/sign-up", data=other)
    token = client.post(
        "/api/auth/log-in",
        data={"email": other["email"], "password": other["password"]},
    ).get_json()["token"]

    response = client.get(
        f"/api/driver-payroll/export-excel/{full_payroll['driver_payroll_code']}",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 404


# ------------------------------------------------------------- other excel exports


@pytest.mark.parametrize(
    "path",
    [
        "/api/drivers/export-excel",
        "/api/routes/export-excel",
        "/api/products/export-excel",
        "/api/shipment-payrolls/export-excel?year=2026",
        "/api/driver-payrolls/export-excel?driver_code=1&start_date=2026-01-01&end_date=2026-12-31",
        "/api/dinatran/export-excel?start_date=2026-01-01&end_date=2026-12-31",
        "/api/statistics/driver/export-excel?start_date=2026-01-01&end_date=2026-12-31",
        "/api/statistics/product/export-excel?start_date=2026-01-01&end_date=2026-12-31",
    ],
)
def test_every_excel_export_answers_a_spreadsheet(client, auth, full_payroll, path):
    response = client.get(path, headers=auth)

    assert response.status_code == 200, response.get_json()
    assert response.data[:2] == b"PK"
